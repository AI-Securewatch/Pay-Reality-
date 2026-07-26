import io

import openpyxl
from docx import Document as DocxDocument

from app.domain.ai_policy_builder.text_extraction import (
    UnsupportedFormatError,
    detect_format,
    extract_text,
)


def test_detect_format_by_extension():
    assert detect_format("memo.pdf", None) == "pdf"
    assert detect_format("grants.docx", None) == "docx"
    assert detect_format("limits.xlsx", None) == "xlsx"
    assert detect_format("limits.csv", None) == "csv"
    assert detect_format("summary.txt", None) == "text"


def test_detect_format_falls_back_to_content_type():
    assert detect_format("upload", "application/pdf") == "pdf"


def test_detect_format_rejects_unknown():
    try:
        detect_format("memo.exe", "application/octet-stream")
        assert False, "expected UnsupportedFormatError"
    except UnsupportedFormatError:
        pass


def test_extract_text_plain_text_has_document_marker():
    result = extract_text("text", b"The Controller may approve up to $50,000.")
    assert "--- document ---" in result
    assert "$50,000" in result


def test_extract_text_csv_has_row_markers():
    result = extract_text("csv", b"role,limit\nController,50000\n")
    assert "--- row 1 ---" in result
    assert "--- row 2 ---" in result
    assert "Controller" in result


def test_extract_text_csv_skips_blank_rows():
    result = extract_text("csv", b"role,limit\n\nController,50000\n")
    assert result.count("---") == 4  # two real rows, blank row skipped


def test_extract_text_docx_has_paragraph_markers():
    doc = DocxDocument()
    doc.add_paragraph("The Regional Controller may approve vendor payments up to $50,000.")
    buf = io.BytesIO()
    doc.save(buf)
    result = extract_text("docx", buf.getvalue())
    assert "--- paragraph 1 ---" in result
    assert "$50,000" in result


def test_extract_text_xlsx_has_sheet_and_row_markers():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Vendors"
    sheet.append(["Role", "Limit"])
    sheet.append(["Regional Controller", 50000])
    buf = io.BytesIO()
    workbook.save(buf)
    result = extract_text("xlsx", buf.getvalue())
    assert "sheet 'Vendors', row 1" in result
    assert "sheet 'Vendors', row 2" in result
    assert "Regional Controller" in result


def test_extract_text_empty_document_is_empty_string_not_an_error():
    assert extract_text("text", b"") == "--- document ---\n"
    assert extract_text("csv", b"") == ""
